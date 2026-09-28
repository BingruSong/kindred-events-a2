"""Create a readable bilingual Word rehearsal guide from the project script."""
import re
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'VIDEO-SCRIPT.md'
OUTPUT = ROOT / 'VIDEO-SCRIPT.docx'
CHINESE_ACTIONS = [
    '通过运行中的 Express 服务打开 http://localhost:3000/ 首页，展示导航和第一组活动卡片。不要直接双击 HTML 文件运行。',
    '打开 api/sql/charityevents.sql，指出建库语句、categories 与 events 表、主键、外键、日期和金额约束及样例数据。导入后运行上方两条 SELECT，展示真实输出；导入失败时不要声称成功。',
    '依次展示 event_db.js、filters.js 和 app.js，指出 SQL 参数占位符与输入校验。请求上方类别、活动列表、组合筛选、单活动和无效日期接口，展示实际 HTTP 状态码和 JSON。',
    '打开浏览器开发者工具的 Network 面板，刷新首页，指出活动请求和响应。打开 clientside/js/app.js，展示 getJSON、initHome、card 和 renderCards 函数。',
    '打开搜索页，选择 Environment，地点输入 Brisbane 并搜索。加入某个结果的实际日期再次搜索，然后点击 Clear Filters。输入不可能存在的地点展示无结果状态，再点击浏览器后退展示条件恢复。',
    '打开一张活动卡片，展示 URL 中的 ID、说明、时间、场地、票价与筹款金额。点击 Register interest 展示弹窗，分别用 Escape 和关闭按钮关闭。打开无效 ID 链接展示错误处理。',
    '运行 npm test 并展示真实输出。简要展示报告的 Data Schema 与 API design 回答，再返回网站；可缩窄浏览器展示响应式布局。结束前确认最终录像不超过 15 分钟。',
]


def add_text(paragraph, text, chinese=False):
    for index, part in enumerate(re.split(r'`([^`]+)`', text)):
        run = paragraph.add_run(part)
        run.font.name = 'Consolas' if index % 2 else 'Arial'
        run.font.size = Pt(11 if index % 2 else 12)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
    paragraph.paragraph_format.line_spacing = 1.3
    paragraph.paragraph_format.space_after = Pt(9)


def block(document, label, text, kind):
    heading = document.add_paragraph()
    heading.paragraph_format.keep_with_next = True
    heading.paragraph_format.space_before = Pt(12)
    run = heading.add_run(label)
    run.bold = True
    run.font.size = Pt(12)
    paragraph = document.add_paragraph()
    add_text(paragraph, text, kind == 'zh')
    if kind == 'action':
        shade = OxmlElement('w:shd')
        shade.set(qn('w:fill'), 'EEF3ED')
        paragraph._p.get_or_add_pPr().append(shade)
    paragraph.paragraph_format.keep_together = True


def main():
    text = SOURCE.read_text(encoding='utf-8')
    sections = re.split(r'^## ', text, flags=re.M)[1:]
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2']:
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), 'Microsoft YaHei')
    for style in document.styles:
        for border in style._element.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    document.styles['Normal'].font.size = Pt(12)
    document.styles['Title'].font.size = Pt(22)
    document.styles['Heading 1'].font.size = Pt(17)
    title = document.add_paragraph('Kindred Events demonstration script', style='Title')
    title.paragraph_format.space_after = Pt(6)
    subtitle = document.add_paragraph('中英双语演讲稿与屏幕操作提醒')
    subtitle.paragraph_format.space_after = Pt(14)
    add_text(document.add_paragraph(),
        'Plan for about 12 minutes 30 seconds including screen actions. Rehearse with a timer and keep the final recording under 15 minutes. Speak the English narration; use the Chinese translation to practise and understand it. The time ranges are a budget, not a measured recording duration.')
    add_text(document.add_paragraph(),
        '按约 12 分 30 秒安排口述与屏幕操作。请计时演练，最终视频不得超过 15 分钟。英文用于口述，中文用于理解和练习；各段时间是预算，不是实测录制时长。录屏时避免显示数据库密码。')
    for index, content in enumerate(sections):
        document.add_page_break()
        heading, body = content.split('\n', 1)
        document.add_paragraph(heading, style='Heading 1')
        action = re.search(r'\*\*Screen action:\*\* (.*?)(?=\n\n\*\*)', body, re.S).group(1).strip()
        english = re.search(r'\*\*Speak \(English\):\*\* (.*?)(?=\n\n\*\*)', body, re.S).group(1).strip()
        chinese = re.search(r'\*\*中文理解：\*\* (.*)', body, re.S).group(1).strip()
        block(document, 'Screen actions  操作提醒 不朗读', action + '\n\n' + CHINESE_ACTIONS[index], 'action')
        block(document, 'English narration  英文口述', english, 'en')
        block(document, 'Chinese translation  中文对照', chinese, 'zh')
    document.core_properties.author = ''
    document.core_properties.last_modified_by = ''
    document.core_properties.title = 'Kindred Events demonstration script'
    for paragraph in document.paragraphs:
        for border in paragraph._p.xpath('./w:pPr/w:pBdr'):
            border.getparent().remove(border)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
