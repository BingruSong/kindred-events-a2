"""Fill a copy of the supplied A2 template without changing its questions."""

from pathlib import Path
from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Pt

SOURCE = Path(r"E:\下载\PROG2002 A2 Report.docx")
OUTPUT = Path(__file__).resolve().parents[1] / "PROG2002 A2 Report.docx"


def add_after(anchor, content):
    node = OxmlElement("w:p")
    anchor._p.addnext(node)
    paragraph = Paragraph(node, anchor._parent)
    paragraph.style = "Normal"
    paragraph.paragraph_format.line_spacing = 1.5
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run(content)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    return paragraph


def main():
    document = Document(SOURCE)
    original = list(document.paragraphs)
    original[5].add_run("Kindred Events")
    for index in (7, 9, 11, 12, 14, 15, 17, 18, 20, 21, 22, 23):
        original[index].paragraph_format.keep_with_next = True

    add_after(original[7],
        "Kindred Events is a fictional charity event discovery website for people who want to find local ways to support a cause. "
        "Visitors can see forthcoming activities, compare causes and locations, and understand each event's purpose and fundraising progress. "
        "The project demonstrates how a browser, a read-only API and a relational database work together to keep event information consistent.")
    add_after(original[9],
        "A visitor needs one clear place to find current and future charity activities and inspect reliable details before deciding whether to attend. "
        "Static event cards would become stale and would not support combined searches. This A2 implementation therefore covers discovery and information display; "
        "registration, payment and administration are outside the implemented scope. The Register control displays the required construction notice.")
    add_after(original[12],
        "MySQL stores categories and events. A Node.js/Express server reads those records through parameterised queries and returns JSON. "
        "The browser uses fetch Promises to request upcoming events, categories, search results and one selected event; JavaScript then creates DOM elements from the response. "
        "The server also serves the three HTML pages and their CSS, JavaScript and original SVG illustrations. Dates are stored as UTC instants and displayed in Australia/Brisbane time.")
    add_after(original[15],
        "The same navigation appears on the home, search and detail pages. Home cards lead directly to details, while search provides labelled date, place and cause controls. "
        "Filters can be combined; Clear Filters resets the controls, URL and results together. Loading, empty and request-failure states use visible messages, with retry on failures. "
        "The selected filters remain in the search URL, so browser back and forward restore the matching state. Responsive layouts, keyboard focus styles, a skip link and an accessible registration dialog support different devices and input methods. "
        "These choices were implemented directly; no user research or wireframe study is claimed.")
    add_after(original[18],
        "The categories table has id (primary key), name (unique) and description. The events table has id (primary key) and category_id (foreign key), making one category relate to many events. "
        "An event stores name, UTC start/end, venue, city, purpose, full description, ticket price, fundraising goal, amount raised, status and illustration path. "
        "DECIMAL columns represent money; constraints reject invalid dates or negative values. The SQL import provides ten active fictional events in five categories and one draft event, which public API queries exclude.")
    add_after(original[21],
        "GET /api/events returns active event summaries and details; optional upcoming=1, date=YYYY-MM-DD, location and category query parameters can be combined with AND. "
        "GET /api/categories returns the cause options. GET /api/events/:id returns the selected active event or 404. "
        "The home page uses /api/events?upcoming=1; search uses that collection endpoint with selected filters.")
    add_after(original[22],
        "Example: GET /api/categories lists available causes for the search control. It expects no request body or path parameter. "
        "A successful response is a JSON array; one seeded element is {\"id\":1,\"name\":\"Environment\",\"description\":\"Hands-on projects that restore and protect local places.\"}. "
        "The browser reads the id as the option value and the name as its label. On a database error the server returns HTTP 500 with an error object, rather than an empty list.")
    add_after(original[23],
        "All three endpoints use GET because they retrieve data without changing it. /api/events is a collection resource, /api/categories is a lookup collection, and /api/events/:id identifies one event. "
        "Invalid query values receive HTTP 400; a missing or unpublished event receives HTTP 404. POST, PUT and DELETE are not implemented because this assignment's specified registration interaction is a construction notice and the site has no data-writing function.")
    original[26].paragraph_format.page_break_before = True
    original[26].paragraph_format.keep_with_next = True
    original[27]._element.getparent().remove(original[27]._element)
    original[30]._element.getparent().remove(original[30]._element)
    original[29]._element.getparent().remove(original[29]._element)
    original[28].text = (
        "I acknowledge that I have used GenAI tools to complete this assessment. I used OpenAI Codex to analyse the assessment requirements, "
        "help design and implement the MySQL schema, Express API and browser interface, generate original SVG illustrations, "
        "assist with tests and troubleshooting, and draft the report and bilingual demonstration script. "
        "This statement records the actual assistance used. The applicable GenAI permission must be confirmed against the course policy before submission."
    )
    original[28].paragraph_format.line_spacing = 1.5
    original[28].paragraph_format.keep_together = True
    for run in original[28].runs:
        run.font.name = "Arial"
        run.font.size = Pt(12)
    anchor = add_after(original[28],
        "Chat log scope: These are selected request excerpts and summaries of AI assistance, not a complete conversation export. "
        "The brief's GenAI Use Level is blank; no assessor approval is claimed.")
    entries = [
        "27 September 2026 — Request excerpt: “开始做这个项目” and “你可以调用智能体辅助你工作。建立计划，一步一步来。” "
        "Codex analysed the brief and planned the project. Delegated agents assisted with requirements, database/API code and frontend code. "
        "Codex generated tests and reviewed the outputs.",
        "28 September 2026 — Request excerpt: “链接BingruSong的GitHub账号 接着完成任务”. Codex created and pushed the repository. "
        "After the explicit request “你直接设置为公开吧，谁都可以访问”, it explained the conflict with the brief's visibility requirement and made the repository public.",
        "28 September 2026 — Request excerpts: “包括视频演讲稿（中英双语带操作提醒）” and “演讲稿做成word，重新放进去”. "
        "Codex drafted and formatted the bilingual Word script with timing budgets and screen actions. It did not record or upload a video.",
        "28 September 2026 — Request excerpt: “按照这个新模板来写，写好后替换掉源文件里的word”. "
        "Codex filled this updated template, acknowledged actual AI use and inspected the rendered report. Identity fields remain blank.",
    ]
    for entry in entries:
        anchor = add_after(anchor, entry)
    add_after(anchor,
        "Verification boundary: Route/filter tests and browser checks used a database stub. Live MySQL verification remains pending because "
        "the local environment has no MySQL server or CLI. The student must review the work and supply a full conversation export if the Unit Assessor requires it.")

    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""
    document.core_properties.comments = ""
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
